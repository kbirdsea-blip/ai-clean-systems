# -*- coding: utf-8 -*-
"""
仲町台店 週次コンサルティングレポート ジェネレーター
----------------------------------------------------
JSON のレポートデータを読み込み、モダンにビジュアル化した
「報告書＋提案書」の Word(.docx) を生成する。

使い方:
    python build_report.py <input.json> <output.docx>

JSON スキーマは SKILL.md および同梱の sample_report.json を参照。
"""
import sys
import json
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ===== カラーパレット =====
NAVY   = '1F3B5C'
BLUE   = '2E5E8C'
GREEN  = '2E7D32'
LGREEN = 'E6F4EA'
LGRAY  = 'EEF1F5'
LBLUE  = 'DDE8F3'
ORANGE = 'E8821E'
LORANGE= 'FBEBD7'
RED    = 'C0392B'
WHITE  = 'FFFFFF'
LINE   = 'C9D3DF'
GREYTX = '666666'

FONT = 'Yu Gothic'


def H(hexstr):
    return RGBColor.from_string(hexstr)


def set_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = H(color)


def shade(el, hexcolor):
    pr = el.get_or_add_tcPr() if el.tag.endswith('}tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    pr.append(shd)


def cell_shade(cell, hexcolor):
    shade(cell._tc, hexcolor)


def set_cell_borders(cell, color=LINE, sz=6, edges=('top', 'left', 'bottom', 'right')):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in edges:
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        borders.append(e)
    tcPr.append(borders)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'none')
        borders.append(e)
    tblPr.append(borders)


def cell_margins(cell, top=60, bottom=60, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        e = OxmlElement(f'w:{k}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def p_run(cell, text, size=10.5, bold=False, color=None, align=None,
          space_after=0, space_before=0, first=True):
    p = cell.paragraphs[0] if (first and not cell.paragraphs[0].runs) else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p


def bullet(cell, text, marker='▸', mcolor=GREEN, size=10.5, color=None):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    r = p.add_run(marker + ' ')
    set_font(r, size=size, bold=True, color=mcolor)
    r2 = p.add_run(text)
    set_font(r2, size=size, color=color)
    return p


def build(data, out_path):
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
        s.top_margin = Cm(1.4)
        s.bottom_margin = Cm(1.4)

    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    def spacer(pts=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        set_font(p.add_run(''), size=pts)
        return p

    def part_band(label):
        spacer(6)
        t = doc.add_table(rows=1, cols=1)
        no_borders(t)
        c = t.cell(0, 0)
        cell_shade(c, NAVY)
        cell_margins(c, 60, 60, 160, 160)
        p_run(c, label, size=11.5, bold=True, color=WHITE)

    def section_bar(num, title, color=BLUE):
        spacer(5)
        t = doc.add_table(rows=1, cols=2)
        no_borders(t)
        t.columns[0].width = Cm(1.1)
        t.columns[1].width = Cm(16)
        cnum = t.cell(0, 0)
        cell_shade(cnum, NAVY)
        cell_margins(cnum, 40, 40, 40, 40)
        cnum.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_run(cnum, str(num), size=14, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        ctit = t.cell(0, 1)
        cell_shade(ctit, color)
        cell_margins(ctit, 60, 60, 160, 60)
        ctit.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_run(ctit, title, size=12.5, bold=True, color=WHITE)

    def callout(title, body, bg=LBLUE, accent=NAVY, icon='💡'):
        t = doc.add_table(rows=1, cols=1)
        no_borders(t)
        c = t.cell(0, 0)
        cell_shade(c, bg)
        set_cell_borders(c, color=accent, sz=4)
        cell_margins(c, 100, 100, 160, 160)
        p_run(c, f'{icon} {title}', size=11, bold=True, color=accent)
        if body:
            p_run(c, body, size=10.5, space_before=3)
        return c

    # ============ タイトルバナー ============
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(banner)
    c = banner.cell(0, 0)
    cell_shade(c, NAVY)
    cell_margins(c, top=170, bottom=170, left=200, right=200)
    p_run(c, data.get('client', 'オーナーズコンパス'), size=10.5, color='C7D6E8',
          align=WD_ALIGN_PARAGRAPH.CENTER)
    p_run(c, f"{data['store']}　週次コンサルティングレポート", size=20, bold=True,
          color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2)
    p_run(c, '― 店舗運営の現状報告と改善提案 ―', size=10.5, color='C7D6E8',
          align=WD_ALIGN_PARAGRAPH.CENTER, space_before=3)

    # 情報バー
    info = doc.add_table(rows=1, cols=4)
    no_borders(info)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ('回 数', f"第{data.get('report_no', 1)}回"),
        ('訪問日', data.get('visit_date', '')),
        ('担当', data.get('consultant', '')),
        ('勤務', data.get('schedule', '')),
    ]
    for i, (k, v) in enumerate(meta):
        cell = info.cell(0, i)
        cell_shade(cell, LGRAY)
        set_cell_borders(cell, color=LINE, sz=4)
        cell_margins(cell, 50, 50, 100, 100)
        p_run(cell, k, size=8, bold=True, color=GREYTX, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_run(cell, v, size=10, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER,
              space_before=1)

    # サマリー
    if data.get('summary'):
        spacer(5)
        callout('本日のサマリー', data['summary'], bg=LBLUE, accent=NAVY, icon='📝')

    # ============ 報告パート ============
    part_band('■ 報告パート ― 本日の実施内容と現状')

    # 1. 実施作業
    section_bar(1, '本日の実施作業（売場改善）', color=BLUE)
    done = data.get('done', [])
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cw = [Cm(5.0), Cm(9.0), Cm(3.6)]
    for i, w in enumerate(cw):
        t.columns[i].width = w
    for cell, head in zip(t.rows[0].cells, ['実施項目', '内容', 'ステータス']):
        cell_shade(cell, NAVY)
        set_cell_borders(cell, color=WHITE, sz=4)
        cell_margins(cell, 50, 50, 100, 100)
        p_run(cell, head, size=10, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ri, d in enumerate(done):
        cells = t.add_row().cells
        bg = WHITE if ri % 2 == 0 else 'F2F6FA'
        status = d.get('status', '完了')
        scolor = GREEN if ('完了' in status) else ORANGE
        vals = [d.get('title', ''), d.get('detail', ''), status]
        for ci, (cell, v) in enumerate(zip(cells, vals)):
            cell_shade(cell, bg)
            set_cell_borders(cell, color=LINE, sz=4)
            cell_margins(cell, 55, 55, 110, 110)
            if ci == 0:
                p_run(cell, v, size=10, bold=True, color=NAVY)
            elif ci == 2:
                p_run(cell, v, size=9.5, bold=True, color=scolor,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                p_run(cell, v, size=10)

    # 2. 現状と気づき
    section_bar(2, '店舗の現状と気づき（課題）', color=BLUE)
    issues = data.get('issues', [])
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate([Cm(8.6), Cm(9.0)]):
        t.columns[i].width = w
    for cell, head in zip(t.rows[0].cells, ['🔍 気づき・現状', '➡ 対応の方向性']):
        cell_shade(cell, NAVY)
        set_cell_borders(cell, color=WHITE, sz=4)
        cell_margins(cell, 50, 50, 110, 110)
        p_run(cell, head, size=10, bold=True, color=WHITE)
    for ri, it in enumerate(issues):
        cells = t.add_row().cells
        cell_shade(cells[0], 'FCEFE2' if ri % 2 else LORANGE)
        cell_shade(cells[1], 'F2F6FA' if ri % 2 else LGREEN)
        for ci, cell in enumerate(cells):
            set_cell_borders(cell, color=LINE, sz=4)
            cell_margins(cell, 60, 60, 120, 120)
        p_run(cells[0], it.get('observation', ''), size=10)
        p_run(cells[1], it.get('action', ''), size=10, color='1B5E20')

    # 3. オーナーへの提案・合意
    section_bar(3, 'オーナーへの提案・合意事項', color=BLUE)
    t = doc.add_table(rows=1, cols=1)
    no_borders(t)
    c = t.cell(0, 0)
    cell_shade(c, LGREEN)
    set_cell_borders(c, color='B7DFC2', sz=4)
    cell_margins(c, 90, 90, 160, 150)
    p_run(c, '🤝 本日オーナーと共有・合意した事項', size=10, bold=True, color=GREEN)
    for item in data.get('proposed', []):
        bullet(c, item, marker='✓', mcolor=GREEN)

    # ============ 提案パート ============
    part_band('■ 提案パート ― 翌週アクションと重点改善提案')

    # 4. 翌週アクションプラン
    section_bar(4, '翌週アクションプラン', color=GREEN)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate([Cm(8.8), Cm(8.8)]):
        t.columns[i].width = w
    head_us = t.cell(0, 0)
    head_st = t.cell(0, 1)
    cell_shade(head_us, BLUE)
    cell_shade(head_st, ORANGE)
    for cell, head in [(head_us, f"👤 当方（{data.get('consultant','担当')}）が行うこと"),
                       (head_st, '🏪 店舗にお願いすること')]:
        set_cell_borders(cell, color=WHITE, sz=4)
        cell_margins(cell, 55, 55, 120, 120)
        p_run(cell, head, size=10, bold=True, color=WHITE)
    body = t.add_row().cells
    cell_shade(body[0], 'EEF3F8')
    cell_shade(body[1], LORANGE)
    for ci, (cell, items, mc) in enumerate([
            (body[0], data.get('next_us', []), BLUE),
            (body[1], data.get('next_store', []), ORANGE)]):
        set_cell_borders(cell, color=LINE, sz=4)
        cell_margins(cell, 80, 80, 130, 130)
        first = True
        for it in items:
            if first:
                p_run(cell, '', size=2)
                first = False
            bullet(cell, it, marker='□', mcolor=mc)

    # 5. 重点課題と改善提案
    recs = data.get('recommendations', [])
    if recs:
        section_bar(5, '重点課題と改善提案', color=ORANGE)
        for idx, rec in enumerate(recs, 1):
            spacer(3)
            tb = doc.add_table(rows=1, cols=1)
            no_borders(tb)
            ct = tb.cell(0, 0)
            cell_shade(ct, BLUE)
            cell_margins(ct, 50, 50, 140, 120)
            p_run(ct, f"提案{idx}　{rec.get('theme','')}", size=11, bold=True, color=WHITE)
            tb2 = doc.add_table(rows=0, cols=2)
            tb2.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, w in enumerate([Cm(2.6), Cm(15.0)]):
                tb2.columns[i].width = w
            rowdefs = [('課題', rec.get('issue', ''), 'FDECEA', RED),
                       ('原因', rec.get('cause', ''), LGRAY, GREYTX),
                       ('対策', rec.get('measure', ''), LGREEN, '1B5E20')]
            for label, val, bg, tc in rowdefs:
                if not val:
                    continue
                cells = tb2.add_row().cells
                cell_shade(cells[0], bg)
                cell_shade(cells[1], WHITE)
                for cell in cells:
                    set_cell_borders(cell, color=LINE, sz=4)
                    cell_margins(cell, 55, 55, 110, 110)
                cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p_run(cells[0], label, size=9.5, bold=True, color=tc,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
                if isinstance(val, list):
                    f = True
                    for v in val:
                        if f:
                            p_run(cells[1], '', size=1)
                            f = False
                        bullet(cells[1], v, marker='▸', mcolor=GREEN, size=10)
                else:
                    p_run(cells[1], val, size=10)

    # 6. ロードマップ
    roadmap = data.get('roadmap', [])
    if roadmap:
        section_bar(6, '優先順位ロードマップ', color=GREEN)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run('一度に全ては回らない。優先順位をつけて段階的に着手する。'), size=10.5)
        n = len(roadmap)
        cols = n * 2 - 1
        road = doc.add_table(rows=1, cols=cols)
        no_borders(road)
        road.alignment = WD_TABLE_ALIGNMENT.CENTER
        stepw = Cm(min(3.8, 15.0 / n))
        for i in range(cols):
            road.columns[i].width = stepw if i % 2 == 0 else Cm(0.7)
        palette = [GREEN, BLUE, BLUE, NAVY, ORANGE, RED]
        si = 0
        for i in range(cols):
            cell = road.cell(0, i)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i % 2 == 0:
                cell_shade(cell, palette[si % len(palette)])
                cell_margins(cell, 70, 70, 40, 40)
                for j, ln in enumerate(roadmap[si].split('\n')):
                    p_run(cell, ln, size=(9 if j == 0 else 10.5), bold=True, color=WHITE,
                          align=WD_ALIGN_PARAGRAPH.CENTER, first=(j == 0))
                si += 1
            else:
                p_run(cell, '➜', size=15, bold=True, color='999999',
                      align=WD_ALIGN_PARAGRAPH.CENTER)

    # 7. KPI
    kpi = data.get('kpi', [])
    if kpi:
        section_bar(7, 'KPI・進捗管理表', color=GREEN)
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hw = [Cm(2.8), Cm(7.6), Cm(3.0), Cm(3.2)]
        for i, w in enumerate(hw):
            table.columns[i].width = w
        for cell, h in zip(table.rows[0].cells, ['項目', '指標（KPI）', '現状', '目標']):
            cell_shade(cell, NAVY)
            set_cell_borders(cell, color=WHITE, sz=4)
            cell_margins(cell, 50, 50, 100, 100)
            p_run(cell, h, size=10, bold=True, color=WHITE)
        for ri, rd in enumerate(kpi):
            cells = table.add_row().cells
            bg = WHITE if ri % 2 == 0 else 'F2F6FA'
            for ci, (cell, v) in enumerate(zip(cells, rd)):
                cell_shade(cell, bg)
                set_cell_borders(cell, color=LINE, sz=4)
                cell_margins(cell, 50, 50, 100, 100)
                p_run(cell, v, size=10, bold=(ci == 0), color=(NAVY if ci == 0 else None))

    # フッター
    if data.get('footer'):
        spacer(6)
        callout('運用のポイント', data['footer'], bg=LBLUE, accent=NAVY, icon='💡')

    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 3:
        print('usage: python build_report.py <input.json> <output.docx>')
        sys.exit(1)
    with open(sys.argv[1], encoding='utf-8') as f:
        data = json.load(f)
    out = build(data, sys.argv[2])
    print('saved:', out)


if __name__ == '__main__':
    main()
