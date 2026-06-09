# -*- coding: utf-8 -*-
"""
オーナーズコンパス 店舗横断サマリー ジェネレーター
--------------------------------------------------
reports/<店舗>/ 配下の各店舗の最新回 JSON を集約し、
全店舗の状況を1枚で見渡せるサマリー Word(.docx) を生成する。

使い方:
    python build_summary.py <reports_dir> <output.docx>
    例) python build_summary.py reports reports/_店舗横断サマリー.docx
"""
import sys
import os
import glob
import json
import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = '1F3B5C'; BLUE = '2E5E8C'; GREEN = '2E7D32'; ORANGE = 'E8821E'
WHITE = 'FFFFFF'; LINE = 'C9D3DF'; LGRAY = 'EEF1F5'; LBLUE = 'DDE8F3'
FONT = 'Yu Gothic'


def H(h):
    return RGBColor.from_string(h)


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = H(color)


def shade(cell, hexcolor):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    pr.append(shd)


def borders(cell, color=LINE, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tcPr.append(b)


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'none'); b.append(e)
    tblPr.append(b)


def margins(cell, t=55, b=55, l=110, r=110):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', t), ('bottom', b), ('start', l), ('end', r)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def put(cell, text, size=10, bold=False, color=None, align=None):
    p = cell.paragraphs[0] if not cell.paragraphs[0].runs else cell.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run(text), size=size, bold=bold, color=color)


def latest_per_store(reports_dir):
    stores = {}
    for store_dir in sorted(glob.glob(os.path.join(reports_dir, '*'))):
        if not os.path.isdir(store_dir):
            continue
        best, best_no = None, -1
        for fp in glob.glob(os.path.join(store_dir, '*.json')):
            try:
                with open(fp, encoding='utf-8') as f:
                    d = json.load(f)
            except Exception:
                continue
            n = d.get('report_no', 0)
            if isinstance(n, int) and n > best_no:
                best, best_no = d, n
        if best:
            stores[os.path.basename(store_dir)] = best
    return stores


def build(reports_dir, out_path):
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(1.6); s.right_margin = Cm(1.6)
        s.top_margin = Cm(1.4); s.bottom_margin = Cm(1.4)
    style = doc.styles['Normal']
    style.font.name = FONT; style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    banner = doc.add_table(rows=1, cols=1); no_borders(banner)
    c = banner.cell(0, 0); shade(c, NAVY); margins(c, 150, 150, 200, 200)
    put(c, 'オーナーズコンパス', 11, color='C7D6E8', align=WD_ALIGN_PARAGRAPH.CENTER)
    put(c, '店舗横断サマリー', 20, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    today = datetime.date.today().strftime('%Y年%m月%d日')
    put(c, f'作成日：{today}', 9.5, color='C7D6E8', align=WD_ALIGN_PARAGRAPH.CENTER)

    stores = latest_per_store(reports_dir)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    set_font(p.add_run(f'対象店舗：{len(stores)}店舗　／　各店の最新回レポートを集約'), size=10, color='555555')

    heads = ['店舗', '最新回', '訪問日', '重点課題（テーマ）', '翌週の主アクション']
    hw = [Cm(3.0), Cm(1.6), Cm(2.6), Cm(5.4), Cm(5.6)]
    table = doc.add_table(rows=1, cols=len(heads))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate(hw):
        table.columns[i].width = w
    for cell, h in zip(table.rows[0].cells, heads):
        shade(cell, NAVY); borders(cell, WHITE); margins(cell)
        put(cell, h, 10, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    for ri, (store, d) in enumerate(sorted(stores.items())):
        cells = table.add_row().cells
        bg = WHITE if ri % 2 == 0 else 'F2F6FA'
        for cell in cells:
            shade(cell, bg); borders(cell); margins(cell)
        cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        put(cells[0], store, 10.5, bold=True, color=NAVY)
        put(cells[1], f"第{d.get('report_no','')}回", 10, align=WD_ALIGN_PARAGRAPH.CENTER)
        put(cells[2], d.get('visit_date', ''), 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        themes = [r.get('theme', '') for r in d.get('recommendations', [])][:4]
        if themes:
            for th in themes:
                pp = cells[3].add_paragraph() if cells[3].paragraphs[0].runs else cells[3].paragraphs[0]
                pp.paragraph_format.space_after = Pt(1)
                set_font(pp.add_run('• '), size=9.5, bold=True, color=ORANGE)
                set_font(pp.add_run(th), size=9.5)
        else:
            put(cells[3], '―', 9.5)
        acts = d.get('next_us', [])[:3]
        if acts:
            for a in acts:
                pp = cells[4].add_paragraph() if cells[4].paragraphs[0].runs else cells[4].paragraphs[0]
                pp.paragraph_format.space_after = Pt(1)
                set_font(pp.add_run('□ '), size=9.5, bold=True, color=BLUE)
                set_font(pp.add_run(a), size=9.5)
        else:
            put(cells[4], '―', 9.5)

    # 各店サマリー文の一覧
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    set_font(p.add_run('■ 各店の総括'), size=12, bold=True, color=NAVY)
    for store, d in sorted(stores.items()):
        if not d.get('summary'):
            continue
        tb = doc.add_table(rows=1, cols=1); no_borders(tb)
        cc = tb.cell(0, 0); shade(cc, LBLUE); borders(cc, NAVY, 4); margins(cc, 80, 80, 150, 150)
        put(cc, f"{store}（第{d.get('report_no','')}回）", 10, bold=True, color=NAVY)
        put(cc, d['summary'], 10)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    doc.save(out_path)
    return out_path, len(stores)


def main():
    reports_dir = sys.argv[1] if len(sys.argv) > 1 else 'reports'
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(reports_dir, '_店舗横断サマリー.docx')
    path, n = build(reports_dir, out)
    print('saved:', path, '| stores:', n)


if __name__ == '__main__':
    main()
