# -*- coding: utf-8 -*-
"""
オーナーズコンパス 店舗 週次コンサルティングレポート ジェネレーター（v2）
------------------------------------------------------------------
JSON のレポートデータを読み込み、モダンにビジュアル化した
「報告書＋提案書」の Word(.docx) を生成する。

v2 の追加機能:
  - 表紙ページ＋目次＋ページ番号（提出書類としての体裁）
  - 週次の進捗トラッキング（同フォルダの過去回 JSON を自動で読み込み、
    「前回アクションの振り返り」「継続課題」「KPI 前回比」を表示）
  - 数値 KPI の達成度ゲージ（棒グラフ風、外部ライブラリ不要）
  - 複数店舗対応（store ごとに reports/<店舗>/ で管理）

使い方:
    python build_report.py <input.json> <output.docx>
"""
import sys
import os
import re
import glob
import json
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ===== カラーパレット =====
NAVY   = '1F3B5C'
BLUE   = '2E5E8C'
GREEN  = '2E7D32'
LGREEN = 'E6F4EA'
LGRAY  = 'EEF1F5'
LBLUE  = 'DDE8F3'
ORANGE = 'E8821E'
LORANGE= 'FBEBD7'
RED    = 'C0392B'
WHITE  = 'FFFFFF'
LINE   = 'C9D3DF'
GREYTX = '666666'
GAUGEBG= 'E3E9F0'

FONT = 'Yu Gothic'


def H(hexstr):
    return RGBColor.from_string(hexstr)


def set_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = H(color)


def shade(el, hexcolor):
    pr = el.get_or_add_tcPr() if el.tag.endswith('}tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    pr.append(shd)


def cell_shade(cell, hexcolor):
    shade(cell._tc, hexcolor)


def set_cell_borders(cell, color=LINE, sz=6, edges=('top', 'left', 'bottom', 'right')):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in edges:
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        borders.append(e)
    tcPr.append(borders)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'none')
        borders.append(e)
    tblPr.append(borders)


def cell_margins(cell, top=60, bottom=60, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        e = OxmlElement(f'w:{k}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def p_run(cell, text, size=10.5, bold=False, color=None, align=None,
          space_after=0, space_before=0, first=True):
    p = cell.paragraphs[0] if (first and not cell.paragraphs[0].runs) else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p


def bullet(cell, text, marker='▸', mcolor=GREEN, size=10.5, color=None):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    r = p.add_run(marker + ' ')
    set_font(r, size=size, bold=True, color=mcolor)
    r2 = p.add_run(text)
    set_font(r2, size=size, color=color)
    return p


def parse_num(v):
    """文字列から数値を取り出す。取れなければ None。"""
    if isinstance(v, (int, float)):
        return float(v)
    if not isinstance(v, str):
        return None
    m = re.search(r'-?\d+(?:\.\d+)?', v.replace(',', ''))
    return float(m.group()) if m else None


def load_previous(input_path, report_no):
    """同じフォルダから report_no が小さい直近の JSON を返す。"""
    folder = os.path.dirname(os.path.abspath(input_path))
    best, best_no = None, -1
    for fp in glob.glob(os.path.join(folder, '*.json')):
        if os.path.abspath(fp) == os.path.abspath(input_path):
            continue
        try:
            with open(fp, encoding='utf-8') as f:
                d = json.load(f)
        except Exception:
            continue
        n = d.get('report_no', 0)
        if isinstance(n, int) and n < report_no and n > best_no:
            best, best_no = d, n
    return best


def build(data, out_path, prev=None):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)

    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    secno = [0]

    def spacer(pts=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        set_font(p.add_run(''), size=pts)
        return p

    def thin_rule(color=NAVY, width=Cm(17.4)):
        t = doc.add_table(rows=1, cols=1)
        no_borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.columns[0].width = width
        c = t.cell(0, 0)
        cell_shade(c, color)
        cell_margins(c, 20, 20, 0, 0)
        p_run(c, '', size=1)

    def part_band(label):
        spacer(6)
        t = doc.add_table(rows=1, cols=1)
        no_borders(t)
        c = t.cell(0, 0)
        cell_shade(c, NAVY)
        cell_margins(c, 60, 60, 160, 160)
        p_run(c, label, size=11.5, bold=True, color=WHITE)

    def section_bar(title, color=BLUE):
        secno[0] += 1
        spacer(5)
        t = doc.add_table(rows=1, cols=2)
        no_borders(t)
        t.columns[0].width = Cm(1.1)
        t.columns[1].width = Cm(16)
        cnum = t.cell(0, 0)
        cell_shade(cnum, NAVY)
        cell_margins(cnum, 40, 40, 40, 40)
        cnum.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_run(cnum, str(secno[0]), size=14, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        ctit = t.cell(0, 1)
        cell_shade(ctit, color)
        cell_margins(ctit, 60, 60, 160, 60)
        ctit.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_run(ctit, title, size=12.5, bold=True, color=WHITE)

    def callout(title, body, bg=LBLUE, accent=NAVY, icon='💡'):
        t = doc.add_table(rows=1, cols=1)
        no_borders(t)
        c = t.cell(0, 0)
        cell_shade(c, bg)
        set_cell_borders(c, color=accent, sz=4)
        cell_margins(c, 100, 100, 160, 160)
        p_run(c, f'{icon} {title}', size=11, bold=True, color=accent)
        if body:
            p_run(c, body, size=10.5, space_before=3)
        return c

    def add_page_number(paragraph):
        run = paragraph.add_run()
        f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE '
        f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
        run._r.append(f1); run._r.append(it); run._r.append(f2)
        set_font(run, size=9, color=GREYTX)

    # ----- フッター（ページ番号） -----
    foot_p = sec.footer.paragraphs[0]
    foot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot_p.add_run(f"{data.get('client','オーナーズコンパス')}｜{data['store']} 週次レポート 第{data.get('report_no',1)}回　　-  ")
    set_font(fr, size=9, color=GREYTX)
    add_page_number(foot_p)
    fr2 = foot_p.add_run('  -')
    set_font(fr2, size=9, color=GREYTX)

    # 目次（存在するセクションを動的に）
    toc = []
    if prev:
        toc.append('前回アクションの振り返り・継続課題')
    toc += ['本日の実施作業（売場改善）']
    if data.get('metrics'):
        toc.append('数値で見る現状（POSデータ）')
    toc += ['店舗の現状と気づき（課題）',
            'オーナーへの提案・合意事項', '翌週アクションプラン']
    if data.get('recommendations'):
        toc.append('重点課題と改善提案')
    if data.get('roadmap'):
        toc.append('優先順位ロードマップ')
    if data.get('kpi'):
        toc.append('KPI・進捗管理表')

    # =========================================================
    # 表紙ページ
    # =========================================================
    for _ in range(3):
        spacer(8)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(data.get('client', 'オーナーズコンパス')), size=13, bold=True, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    set_font(p.add_run('店舗コンサルティング'), size=10.5, color=GREYTX)
    spacer(10)
    thin_rule(NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    set_font(p.add_run(data['store']), size=34, bold=True, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run('週次コンサルティングレポート'), size=18, bold=True, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    set_font(p.add_run('― 店舗運営の現状報告と改善提案 ―'), size=11, color=GREYTX)
    spacer(6)
    thin_rule(NAVY)

    spacer(14)
    meta = doc.add_table(rows=1, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.columns[0].width = Cm(3.2)
    meta.columns[1].width = Cm(8.0)
    rows = [
        ('回　数', f"第{data.get('report_no',1)}回"),
        ('訪問日', data.get('visit_date', '')),
        ('担　当', data.get('consultant', '')),
        ('勤　務', data.get('schedule', '')),
        ('作成日', data.get('created', data.get('visit_date', ''))),
    ]
    first = True
    for k, v in rows:
        cells = (meta.rows[0].cells if first else meta.add_row().cells)
        first = False
        cell_shade(cells[0], NAVY)
        cell_shade(cells[1], LGRAY)
        for c in cells:
            set_cell_borders(c, color=WHITE, sz=6)
            cell_margins(c, 60, 60, 130, 130)
        p_run(cells[0], k, size=10, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_run(cells[1], v, size=11, bold=True, color=NAVY)

    spacer(16)
    tb = doc.add_table(rows=1, cols=1)
    no_borders(tb)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.columns[0].width = Cm(11.2)
    c = tb.cell(0, 0)
    cell_shade(c, 'F7F9FC')
    set_cell_borders(c, color=LINE, sz=4)
    cell_margins(c, 110, 120, 180, 180)
    p_run(c, '目　次', size=12, bold=True, color=NAVY)
    for i, t in enumerate(toc, 1):
        pp = c.add_paragraph(); pp.paragraph_format.space_after = Pt(3)
        pp.paragraph_format.space_before = Pt(2)
        r = pp.add_run(f'{i}.　'); set_font(r, size=10.5, bold=True, color=BLUE)
        r2 = pp.add_run(t); set_font(r2, size=10.5, color='333333')

    doc.add_page_break()

    # =========================================================
    # 本文：タイトルバナー
    # =========================================================
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(banner)
    c = banner.cell(0, 0)
    cell_shade(c, NAVY)
    cell_margins(c, top=150, bottom=150, left=200, right=200)
    p_run(c, f"{data['store']}　週次コンサルティングレポート", size=18, bold=True,
          color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_run(c, f"第{data.get('report_no',1)}回　／　{data.get('visit_date','')}　／　担当：{data.get('consultant','')}",
          size=10, color='C7D6E8', align=WD_ALIGN_PARAGRAPH.CENTER, space_before=3)

    if data.get('summary'):
        spacer(5)
        callout('本日のサマリー', data['summary'], bg=LBLUE, accent=NAVY, icon='📝')

    # =========================================================
    # 報告パート
    # =========================================================
    part_band('■ 報告パート ― 本日の実施内容と現状')

    if prev:
        section_bar('前回アクションの振り返り・継続課題', color=BLUE)
        followups = data.get('followups')
        if not followups:
            followups = ([{'owner': '当方', 'item': x, 'status': '確認中'} for x in prev.get('next_us', [])]
                         + [{'owner': '店舗', 'item': x, 'status': '確認中'} for x in prev.get('next_store', [])])
        t = doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, w in enumerate([Cm(2.4), Cm(11.4), Cm(3.6)]):
            t.columns[i].width = w
        for cell, head in zip(t.rows[0].cells, ['区分', f'前回（第{prev.get("report_no","")}回）のアクション', '状況']):
            cell_shade(cell, NAVY); set_cell_borders(cell, color=WHITE, sz=4)
            cell_margins(cell, 50, 50, 100, 100)
            p_run(cell, head, size=10, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        for ri, fu in enumerate(followups):
            cells = t.add_row().cells
            bg = WHITE if ri % 2 == 0 else 'F2F6FA'
            st = fu.get('status', '確認中')
            sc = GREEN if '完了' in st else (ORANGE if ('進行' in st or '一部' in st) else GREYTX)
            owner = fu.get('owner', '')
            oc = BLUE if owner == '当方' else ORANGE
            for cell in cells:
                cell_shade(cell, bg); set_cell_borders(cell, color=LINE, sz=4)
                cell_margins(cell, 55, 55, 110, 110)
            p_run(cells[0], owner, size=9.5, bold=True, color=oc, align=WD_ALIGN_PARAGRAPH.CENTER)
            txt = fu.get('item', '')
            if fu.get('note'):
                txt += f"（{fu['note']}）"
            p_run(cells[1], txt, size=10)
            p_run(cells[2], st, size=9.5, bold=True, color=sc, align=WD_ALIGN_PARAGRAPH.CENTER)
        carry = data.get('carryover', [])
        if carry:
            t2 = doc.add_table(rows=1, cols=1); no_borders(t2)
            cc = t2.cell(0, 0); cell_shade(cc, LORANGE); set_cell_borders(cc, color=ORANGE, sz=4)
            cell_margins(cc, 80, 80, 160, 150)
            p_run(cc, '⚠ 継続課題（引き続き取り組む）', size=10, bold=True, color=ORANGE)
            for x in carry:
                bullet(cc, x, marker='・', mcolor=ORANGE)

    # 実施作業
    section_bar('本日の実施作業（売場改善）', color=BLUE)
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate([Cm(5.0), Cm(9.0), Cm(3.6)]):
        t.columns[i].width = w
    for cell, head in zip(t.rows[0].cells, ['実施項目', '内容', 'ステータス']):
        cell_shade(cell, NAVY); set_cell_borders(cell, color=WHITE, sz=4)
        cell_margins(cell, 50, 50, 100, 100)
        p_run(cell, head, size=10, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ri, d in enumerate(data.get('done', [])):
        cells = t.add_row().cells
        bg = WHITE if ri % 2 == 0 else 'F2F6FA'
        status = d.get('status', '完了')
        scolor = GREEN if ('完了' in status) else ORANGE
        for ci, (cell, v) in enumerate(zip(cells, [d.get('title', ''), d.get('detail', ''), status])):
            cell_shade(cell, bg); set_cell_borders(cell, color=LINE, sz=4)
            cell_margins(cell, 55, 55, 110, 110)
            if ci == 0:
                p_run(cell, v, size=10, bold=True, color=NAVY)
            elif ci == 2:
                p_run(cell, v, size=9.5, bold=True, color=scolor, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                p_run(cell, v, size=10)

    # 数値で見る現状（POSデータ）
    metrics = data.get('metrics', [])
    if metrics:
        section_bar('数値で見る現状（POSデータ）', color=BLUE)
        if data.get('metrics_period'):
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
            set_font(p.add_run(f"出典：{data['metrics_period']}"), size=9, color=GREYTX)
        heads = ['分類', '指標', '自店', '基準', '読み取り']
        hw = [Cm(1.8), Cm(3.4), Cm(3.4), Cm(2.6), Cm(6.4)]
        mt = doc.add_table(rows=1, cols=len(heads))
        mt.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, w in enumerate(hw):
            mt.columns[i].width = w
        for cell, h in zip(mt.rows[0].cells, heads):
            cell_shade(cell, NAVY); set_cell_borders(cell, color=WHITE, sz=4)
            cell_margins(cell, 50, 50, 90, 90)
            p_run(cell, h, size=10, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        for ri, rd in enumerate(metrics):
            row = list(rd) + [''] * (len(heads) - len(rd))
            cells = mt.add_row().cells
            bg = WHITE if ri % 2 == 0 else 'F2F6FA'
            for ci, (cell, v) in enumerate(zip(cells, row)):
                cell_shade(cell, bg); set_cell_borders(cell, color=LINE, sz=4)
                cell_margins(cell, 50, 50, 90, 90)
                p_run(cell, str(v), size=9.5, bold=(ci == 0), color=(NAVY if ci == 0 else None))
        good = data.get('metrics_good', [])
        warn = data.get('metrics_warn', [])
        if good or warn:
            spacer(3)
            t = doc.add_table(rows=1, cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, w in enumerate([Cm(8.8), Cm(8.8)]):
                t.columns[i].width = w
            cell_shade(t.cell(0, 0), GREEN); cell_shade(t.cell(0, 1), RED)
            for cell, head in [(t.cell(0, 0), '📈 良い兆し'), (t.cell(0, 1), '⚠ 注意シグナル')]:
                set_cell_borders(cell, color=WHITE, sz=4); cell_margins(cell, 55, 55, 120, 120)
                p_run(cell, head, size=10, bold=True, color=WHITE)
            body = t.add_row().cells
            cell_shade(body[0], LGREEN); cell_shade(body[1], 'FDECEA')
            for cell, items, mc in [(body[0], good, GREEN), (body[1], warn, RED)]:
                set_cell_borders(cell, color=LINE, sz=4); cell_margins(cell, 70, 70, 120, 120)
                f = True
                for it in items:
                    if f:
                        p_run(cell, '', size=2); f = False
                    bullet(cell, it, marker='・', mcolor=mc, size=9.5)

    # 現状と気づき
    section_bar('店舗の現状と気づき（課題）', color=BLUE)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate([Cm(8.6), Cm(9.0)]):
        t.columns[i].width = w
    for cell, head in zip(t.rows[0].cells, ['🔍 気づき・現状', '➡ 対応の方向性']):
        cell_shade(cell, NAVY); set_cell_borders(cell, color=WHITE, sz=4)
        cell_margins(cell, 50, 50, 110, 110)
        p_run(cell, head, size=10, bold=True, color=WHITE)
    for ri, it in enumerate(data.get('issues', [])):
        cells = t.add_row().cells
        cell_shade(cells[0], 'FCEFE2' if ri % 2 else LORANGE)
        cell_shade(cells[1], 'F2F6FA' if ri % 2 else LGREEN)
        for cell in cells:
            set_cell_borders(cell, color=LINE, sz=4); cell_margins(cell, 60, 60, 120, 120)
        p_run(cells[0], it.get('observation', ''), size=10)
        p_run(cells[1], it.get('action', ''), size=10, color='1B5E20')

    # オーナー合意
    section_bar('オーナーへの提案・合意事項', color=BLUE)
    t = doc.add_table(rows=1, cols=1); no_borders(t)
    c = t.cell(0, 0); cell_shade(c, LGREEN); set_cell_borders(c, color='B7DFC2', sz=4)
    cell_margins(c, 90, 90, 160, 150)
    p_run(c, '🤝 本日オーナーと共有・合意した事項', size=10, bold=True, color=GREEN)
    for item in data.get('proposed', []):
        bullet(c, item, marker='✓', mcolor=GREEN)

    # =========================================================
    # 提案パート
    # =========================================================
    part_band('■ 提案パート ― 翌週アクションと重点改善提案')

    section_bar('翌週アクションプラン', color=GREEN)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate([Cm(8.8), Cm(8.8)]):
        t.columns[i].width = w
    cell_shade(t.cell(0, 0), BLUE); cell_shade(t.cell(0, 1), ORANGE)
    for cell, head in [(t.cell(0, 0), f"👤 当方（{data.get('consultant','担当')}）が行うこと"),
                       (t.cell(0, 1), '🏪 店舗にお願いすること')]:
        set_cell_borders(cell, color=WHITE, sz=4); cell_margins(cell, 55, 55, 120, 120)
        p_run(cell, head, size=10, bold=True, color=WHITE)
    body = t.add_row().cells
    cell_shade(body[0], 'EEF3F8'); cell_shade(body[1], LORANGE)
    for cell, items, mc in [(body[0], data.get('next_us', []), BLUE),
                            (body[1], data.get('next_store', []), ORANGE)]:
        set_cell_borders(cell, color=LINE, sz=4); cell_margins(cell, 80, 80, 130, 130)
        f = True
        for it in items:
            if f:
                p_run(cell, '', size=2); f = False
            bullet(cell, it, marker='□', mcolor=mc)

    recs = data.get('recommendations', [])
    if recs:
        section_bar('重点課題と改善提案', color=ORANGE)
        for idx, rec in enumerate(recs, 1):
            spacer(3)
            tb = doc.add_table(rows=1, cols=1); no_borders(tb)
            ct = tb.cell(0, 0); cell_shade(ct, BLUE); cell_margins(ct, 50, 50, 140, 120)
            p_run(ct, f"提案{idx}　{rec.get('theme','')}", size=11, bold=True, color=WHITE)
            tb2 = doc.add_table(rows=0, cols=2); tb2.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, w in enumerate([Cm(2.6), Cm(15.0)]):
                tb2.columns[i].width = w
            for label, val, bg, tc in [('課題', rec.get('issue', ''), 'FDECEA', RED),
                                       ('原因', rec.get('cause', ''), LGRAY, GREYTX),
                                       ('対策', rec.get('measure', ''), LGREEN, '1B5E20')]:
                if not val:
                    continue
                cells = tb2.add_row().cells
                cell_shade(cells[0], bg); cell_shade(cells[1], WHITE)
                for cell in cells:
                    set_cell_borders(cell, color=LINE, sz=4); cell_margins(cell, 55, 55, 110, 110)
                cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p_run(cells[0], label, size=9.5, bold=True, color=tc, align=WD_ALIGN_PARAGRAPH.CENTER)
                if isinstance(val, list):
                    f = True
                    for v in val:
                        if f:
                            p_run(cells[1], '', size=1); f = False
                        bullet(cells[1], v, marker='▸', mcolor=GREEN, size=10)
                else:
                    p_run(cells[1], val, size=10)

    roadmap = data.get('roadmap', [])
    if roadmap:
        section_bar('優先順位ロードマップ', color=GREEN)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run('一度に全ては回らない。優先順位をつけて段階的に着手する。'), size=10.5)
        n = len(roadmap)
        cols = n * 2 - 1
        road = doc.add_table(rows=1, cols=cols); no_borders(road)
        road.alignment = WD_TABLE_ALIGNMENT.CENTER
        stepw = Cm(min(3.8, 15.0 / n))
        for i in range(cols):
            road.columns[i].width = stepw if i % 2 == 0 else Cm(0.7)
        palette = [GREEN, BLUE, BLUE, NAVY, ORANGE, RED]
        si = 0
        for i in range(cols):
            cell = road.cell(0, i)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i % 2 == 0:
                cell_shade(cell, palette[si % len(palette)]); cell_margins(cell, 70, 70, 40, 40)
                for j, ln in enumerate(roadmap[si].split('\n')):
                    p_run(cell, ln, size=(9 if j == 0 else 10.5), bold=True, color=WHITE,
                          align=WD_ALIGN_PARAGRAPH.CENTER, first=(j == 0))
                si += 1
            else:
                p_run(cell, '➜', size=15, bold=True, color='999999', align=WD_ALIGN_PARAGRAPH.CENTER)

    # KPI（前回比つき）＋達成度ゲージ
    kpi = data.get('kpi', [])
    if kpi:
        section_bar('KPI・進捗管理表', color=GREEN)
        prev_map = {}
        if prev:
            for r in prev.get('kpi', []):
                if len(r) >= 3:
                    prev_map[(r[0], r[1])] = r[2]
        show_prev = bool(prev_map)
        if show_prev:
            heads = ['項目', '指標（KPI）', '前回', '現状', '目標']
            hw = [Cm(2.6), Cm(6.4), Cm(2.6), Cm(2.6), Cm(3.2)]
        else:
            heads = ['項目', '指標（KPI）', '現状', '目標']
            hw = [Cm(2.8), Cm(7.6), Cm(3.0), Cm(3.2)]
        table = doc.add_table(rows=1, cols=len(heads))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, w in enumerate(hw):
            table.columns[i].width = w
        for cell, h in zip(table.rows[0].cells, heads):
            cell_shade(cell, NAVY); set_cell_borders(cell, color=WHITE, sz=4)
            cell_margins(cell, 50, 50, 100, 100)
            p_run(cell, h, size=10, bold=True, color=WHITE)
        for ri, rd in enumerate(kpi):
            item, ind = rd[0], rd[1]
            cur = rd[2] if len(rd) > 2 else ''
            tgt = rd[3] if len(rd) > 3 else ''
            row = [item, ind] + ([prev_map.get((item, ind), '―')] if show_prev else []) + [cur, tgt]
            cells = table.add_row().cells
            bg = WHITE if ri % 2 == 0 else 'F2F6FA'
            for ci, (cell, v) in enumerate(zip(cells, row)):
                cell_shade(cell, bg); set_cell_borders(cell, color=LINE, sz=4)
                cell_margins(cell, 50, 50, 100, 100)
                p_run(cell, str(v), size=10, bold=(ci == 0), color=(NAVY if ci == 0 else None))

        # 達成度ゲージ（現状・目標が数値のものだけ）。ネストせず1行=1テーブル。
        gauges = []
        for rd in kpi:
            if len(rd) < 4:
                continue
            cv, tv = parse_num(rd[2]), parse_num(rd[3])
            if cv is not None and tv is not None and tv != 0:
                gauges.append((rd[1], min(max(cv / tv, 0.0), 1.0)))
        if gauges:
            spacer(4)
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
            set_font(p.add_run('▍ 数値目標の達成度'), size=10.5, bold=True, color=NAVY)
            BAR = 8.4
            for label, ratio in gauges:
                fill_w = max(0.06, BAR * ratio)
                rem_w = max(0.0, BAR - fill_w)
                gt = doc.add_table(rows=1, cols=4); no_borders(gt)
                gt.alignment = WD_TABLE_ALIGNMENT.CENTER
                gt.columns[0].width = Cm(6.0)
                gt.columns[1].width = Cm(fill_w)
                gt.columns[2].width = Cm(rem_w)
                gt.columns[3].width = Cm(2.0)
                lc = gt.cell(0, 0); lc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p_run(lc, label, size=9.5)
                fc = gt.cell(0, 1); cell_shade(fc, GREEN if ratio >= 1 else BLUE)
                cell_margins(fc, 36, 36, 10, 10); p_run(fc, '', size=6)
                rc = gt.cell(0, 2); cell_shade(rc, GAUGEBG)
                cell_margins(rc, 36, 36, 10, 10); p_run(rc, '', size=6)
                pc = gt.cell(0, 3); pc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p_run(pc, f"{int(round(ratio*100))}%", size=9.5, bold=True,
                      color=(GREEN if ratio >= 1 else BLUE), align=WD_ALIGN_PARAGRAPH.RIGHT)

    if data.get('footer'):
        spacer(6)
        callout('運用のポイント', data['footer'], bg=LBLUE, accent=NAVY, icon='💡')

    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 3:
        print('usage: python build_report.py <input.json> <output.docx>')
        sys.exit(1)
    with open(sys.argv[1], encoding='utf-8') as f:
        data = json.load(f)
    prev = load_previous(sys.argv[1], data.get('report_no', 1))
    out = build(data, sys.argv[2], prev=prev)
    print('saved:', out, '| previous:', (prev.get('report_no') if prev else None))


if __name__ == '__main__':
    main()
