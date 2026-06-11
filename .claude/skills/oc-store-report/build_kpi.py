# -*- coding: utf-8 -*-
"""
オーナーズコンパス 店舗KPIダッシュボード ジェネレーター
------------------------------------------------------
POSサマリーシート等から起こした KPI JSON を読み込み、
1〜2ページの KPIダッシュボード Word(.docx) を生成する。
グラフは外部ライブラリを使わず、塗り分けセルの横棒で描画する。

使い方:
    python build_kpi.py <kpi.json> <output.docx>
"""
import sys
import json
import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = '1F3B5C'; BLUE = '2E5E8C'; GREEN = '2E7D32'; ORANGE = 'E8821E'
RED = 'C0392B'; WHITE = 'FFFFFF'; LINE = 'C9D3DF'; LGRAY = 'EEF1F5'
LBLUE = 'DDE8F3'; LGREEN = 'E6F4EA'; LORANGE = 'FBEBD7'; LRED = 'FDECEA'
GREYTX = '666666'; GAUGEBG = 'E3E9F0'
FONT = 'Yu Gothic'

TONE = {
    'good': (LGREEN, GREEN), 'warn': (LORANGE, ORANGE),
    'bad': (LRED, RED), 'flat': (LGRAY, GREYTX),
}


def H(h):
    return RGBColor.from_string(h)


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = H(color)


def shade(cell, hexcolor):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    pr.append(shd)


def borders(cell, color=LINE, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tcPr.append(b)


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'none'); b.append(e)
    tblPr.append(b)


def margins(cell, t=50, b=50, l=100, r=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', t), ('bottom', b), ('start', l), ('end', r)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def put(cell, text, size=10, bold=False, color=None, align=None, before=0, after=1, first=True):
    p = cell.paragraphs[0] if (first and not cell.paragraphs[0].runs) else cell.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def build(data, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.2); sec.bottom_margin = Cm(1.2)
    style = doc.styles['Normal']
    style.font.name = FONT; style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    def spacer(pts=5):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        set_font(p.add_run(''), size=pts)

    def heading(text, color=NAVY):
        spacer(4)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run('▍ '), size=12, bold=True, color=color)
        set_font(p.add_run(text), size=12, bold=True, color=NAVY)

    # ----- ヘッダーバナー -----
    banner = doc.add_table(rows=1, cols=1); no_borders(banner)
    c = banner.cell(0, 0); shade(c, NAVY); margins(c, 130, 130, 200, 200)
    put(c, data.get('client', 'オーナーズコンパス'), 10, color='C7D6E8', align=WD_ALIGN_PARAGRAPH.CENTER)
    put(c, f"{data['store']}　KPIダッシュボード", 19, bold=True, color=WHITE,
        align=WD_ALIGN_PARAGRAPH.CENTER)
    created = data.get('created', datetime.date.today().strftime('%Y年%m月%d日'))
    put(c, f"対象期間：{data.get('period','')}　／　作成日：{created}", 9.5, color='C7D6E8',
        align=WD_ALIGN_PARAGRAPH.CENTER)
    if data.get('source'):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(f"出典：{data['source']}"), size=8.5, color=GREYTX)

    # ----- ヘッドラインKPIカード -----
    cards = data.get('headline', [])
    if cards:
        heading('今週のヘッドラインKPI')
        for i in range(0, len(cards), 3):
            chunk = cards[i:i + 3]
            t = doc.add_table(rows=1, cols=3)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ci in range(3):
                t.columns[ci].width = Cm(5.9)
            for ci in range(3):
                cell = t.cell(0, ci)
                if ci >= len(chunk):
                    borders(cell, WHITE, 2)
                    continue
                cd = chunk[ci]
                bg, ac = TONE.get(cd.get('tone', 'flat'), TONE['flat'])
                shade(cell, bg); borders(cell, ac, 6); margins(cell, 70, 70, 110, 110)
                put(cell, cd.get('label', ''), 9, bold=True, color=ac)
                put(cell, cd.get('value', ''), 16, bold=True, color=NAVY)
                if cd.get('sub'):
                    put(cell, cd['sub'], 8.5, color='333333')
                if cd.get('base'):
                    put(cell, cd['base'], 8.5, color=GREYTX)
            spacer(2)

    # ----- 横棒チャート部品 -----
    def bar_row(label, vtxt, ratio, fill_color, label_w=1.6, barw=10.2, val_w=4.4):
        fill = max(0.08, barw * ratio)
        rest = max(0.0, barw - fill)
        t = doc.add_table(rows=1, cols=4); no_borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, w in enumerate([Cm(label_w), Cm(fill), Cm(rest), Cm(val_w)]):
            t.columns[i].width = w
        lc = t.cell(0, 0); lc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        put(lc, label, 9, bold=True, color=NAVY)
        fc = t.cell(0, 1); shade(fc, fill_color); margins(fc, 30, 30, 10, 10)
        put(fc, '', 5)
        rc = t.cell(0, 2); shade(rc, GAUGEBG); margins(rc, 30, 30, 10, 10)
        put(rc, '', 5)
        vc = t.cell(0, 3); vc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        put(vc, vtxt, 8.5, color='333333')

    # ----- 曜日別トレンド -----
    series = data.get('daily_series', [])
    if series:
        heading('曜日別トレンド（横棒＝当週実績）')
        for s in series:
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(4)
            set_font(p.add_run(s.get('name', '')), size=10, bold=True, color=BLUE)
            rows = s.get('rows', [])
            vals = [r[1] for r in rows]
            mx = max(vals) if vals else 1
            color = {'red': RED, 'green': GREEN, 'orange': ORANGE}.get(s.get('kind', ''), BLUE)
            for r in rows:
                label, v = r[0], r[1]
                note = f"（{r[2]}）" if len(r) > 2 and r[2] else ''
                bar_row(label, f"{v:,}{note}", (v / mx if mx else 0), color)
            if s.get('note'):
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                set_font(p.add_run(f"※ {s['note']}"), size=8.5, color=GREYTX)

    # ----- カテゴリ別 前年差額 -----
    cat = data.get('category_diff')
    if cat:
        heading(f"カテゴリ別 前年差額（{cat.get('unit','千円')}）　上位／下位")
        ups = cat.get('up', []); downs = cat.get('down', [])
        mx = max([abs(v) for _, v in ups + downs] or [1])
        for name, v in ups:
            bar_row(name, f"＋{v}", abs(v) / mx, GREEN, label_w=4.2, barw=8.2, val_w=2.2)
        for name, v in downs:
            bar_row(name, f"{v}", abs(v) / mx, RED, label_w=4.2, barw=8.2, val_w=2.2)
        if cat.get('note'):
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run(f"※ {cat['note']}"), size=8.5, color=GREYTX)

    # ----- 注目ポイント -----
    hl = data.get('highlights', [])
    if hl:
        heading('数字の読みどころ')
        t = doc.add_table(rows=1, cols=1); no_borders(t)
        c = t.cell(0, 0); shade(c, LBLUE); borders(c, NAVY, 4); margins(c, 90, 90, 150, 150)
        first = True
        for x in hl:
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.4)
            set_font(p.add_run('◆ '), size=9.5, bold=True, color=NAVY)
            set_font(p.add_run(x), size=9.5)

    # ----- KPI×打ち手 -----
    acts = data.get('actions', [])
    if acts:
        heading('KPIシグナル → 今週の打ち手')
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.columns[0].width = Cm(6.2); t.columns[1].width = Cm(11.6)
        for cell, h in zip(t.rows[0].cells, ['KPIシグナル', '打ち手（週次レポートのアクションと連動）']):
            shade(cell, NAVY); borders(cell, WHITE, 4); margins(cell)
            put(cell, h, 9.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        for ri, row in enumerate(acts):
            cells = t.add_row().cells
            bg = WHITE if ri % 2 == 0 else 'F2F6FA'
            for cell in cells:
                shade(cell, bg); borders(cell); margins(cell)
            put(cells[0], row[0], 9.5, bold=True, color=RED)
            put(cells[1], row[1] if len(row) > 1 else '', 9.5)

    if data.get('footer'):
        spacer(4)
        t = doc.add_table(rows=1, cols=1); no_borders(t)
        c = t.cell(0, 0); shade(c, LGRAY); borders(c, LINE, 4); margins(c, 70, 70, 140, 140)
        put(c, f"💡 {data['footer']}", 9.5, color='333333')

    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 3:
        print('usage: python build_kpi.py <kpi.json> <output.docx>')
        sys.exit(1)
    with open(sys.argv[1], encoding='utf-8') as f:
        data = json.load(f)
    print('saved:', build(data, sys.argv[2]))


if __name__ == '__main__':
    main()
